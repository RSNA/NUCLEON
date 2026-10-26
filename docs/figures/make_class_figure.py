#!/usr/bin/env python3
"""Draw Figure 1 (NUCLEON classes and properties) as an editable PowerPoint slide.

The style template is the author's Fig1 slide (template/Fig1_template.pptx): its
header, property-panel, annotation-panel, arrow and label shapes are cloned and
placed for the six upper classes of NUCLEON v1.6. Writes class_figure.pptx; use
LibreOffice to export PDF/PNG (see README).

Usage: python3 docs/figures/make_class_figure.py
"""
import copy
import os

from lxml import etree
from pptx import Presentation

HERE = os.path.dirname(os.path.abspath(__file__))
TEMPLATE = os.path.join(HERE, "template", "Fig1_template.pptx")
OUT = os.path.join(HERE, "class_figure.pptx")
A = "{http://schemas.openxmlformats.org/drawingml/2006/main}"
P = "{http://schemas.openxmlformats.org/presentationml/2006/main}"
IN = 914400

prs = Presentation(TEMPLATE)
slide = prs.slides[0]
tree = slide.shapes._spTree
by_name = {sh.name: sh._element for sh in slide.shapes}
T = {k: copy.deepcopy(by_name[v]) for k, v in {
    "header": "Rectangle 3", "body": "Rectangle 4", "annot": "Rectangle 6",
    "arrow": "Straight Arrow Connector 14", "label": "TextBox 16"}.items()}
for el in list(by_name.values()):
    tree.remove(el)

next_id = [100]


def place(kind, x, y, w, h):
    el = copy.deepcopy(T[kind])
    cnv = el.find(f".//{P}cNvPr")
    cnv.set("id", str(next_id[0]))
    cnv.set("name", f"{kind} {next_id[0]}")
    for ext in cnv.findall(f"{A}extLst"):
        cnv.remove(ext)
    next_id[0] += 1
    xfrm = el.find(f".//{A}xfrm")
    for k in list(xfrm.attrib):
        del xfrm.attrib[k]
    if w < 0:                                   # arrow pointing left
        xfrm.set("flipH", "1"); x, w = x + w, -w
    if h < 0:                                   # arrow pointing up
        xfrm.set("flipV", "1"); y, h = y + h, -h
    xfrm.find(f"{A}off").attrib.update({"x": str(int(x)), "y": str(int(y))})
    xfrm.find(f"{A}ext").attrib.update({"cx": str(int(w)), "cy": str(int(h))})
    tree.append(el)
    return el


def set_lines(el, lines, align=None):
    body = el.find(f".//{P}txBody")
    ps = body.findall(f"{A}p")
    ptpl = ps[0]
    runs = ptpl.findall(f"{A}r")
    r_main = runs[0].find(f"{A}rPr")
    r_arrow = runs[1].find(f"{A}rPr") if len(runs) > 1 else r_main
    end = ps[-1].find(f"{A}endParaRPr")
    for p in ps:
        body.remove(p)
    for line in lines:
        p = etree.SubElement(body, f"{A}p")
        ppr = ptpl.find(f"{A}pPr")
        if ppr is not None:
            p.append(copy.deepcopy(ppr))
        parts = line.split("→", 1)
        segs = [(parts[0], r_main)] + ([("→" + parts[1], r_arrow)] if len(parts) == 2 else [])
        for txt, rpr in segs:
            r = etree.SubElement(p, f"{A}r")
            r.append(copy.deepcopy(rpr))
            etree.SubElement(r, f"{A}t").text = txt
        if end is not None:
            e = copy.deepcopy(end)
            sz = rpr.get("sz")
            if sz:
                e.set("sz", sz)
            p.append(e)


# geometry (EMU)
W = int(2.8 * IN)
GAP = int(1.85 * IN)
X0 = (prs.slide_width - (3 * W + 2 * GAP)) // 2
COL = [X0, X0 + W + GAP, X0 + 2 * (W + GAP)]
Y0 = int(0.45 * IN)
HDR = 457200
LINE14, LINE12, PAD = 226000, 200000, 80000
VGAP = int(0.95 * IN)


def klass(col, y, name, props, annots):
    x = COL[col]
    set_lines(place("header", x, y, W, HDR), [name])
    yy = y + HDR
    if props:
        h = LINE14 * len(props) + PAD
        set_lines(place("body", x, yy, W, h), props)
        yy += h
    if annots:
        n = sum(2 if len(a) > 30 else 1 for a in annots)   # long lines wrap once
        h = LINE12 * n + PAD
        set_lines(place("annot", x, yy, W, h), annots)
        yy += h
    return x, y, yy                              # left, top, bottom


def arrow(x1, y1, x2, y2):
    place("arrow", x1, y1, x2 - x1, y2 - y1)


def label(text, cx, y, align="ctr", w=int(1.9 * IN)):
    el = place("label", cx - w / 2 if align == "ctr" else (cx if align == "l" else cx - w), y, w, 338554)
    el.find(f".//{A}p/{A}pPr").set("algn", align)
    set_lines(el, [text])


el_ = klass(0, Y0, ":Element",
            [":atomicNumber", ":groupNumber", ":periodNumber"],
            ["skos:exactMatch →", "      OCE, SNOMED CT", "rdfs:seeAlso → PubChem", "skos:altLabel (7 languages)"])
nu = klass(1, Y0, ":Nuclide",
           [":atomicNumber", ":massNumber", ":neutronNumber", ":metastableState", ":excitationEnergy",
            ":isStable", ":halfLife", ":naturalAbundance", ":decayWidthValue / -Unit"],
           [":halfLifeValue / -Unit", ":spinParity", ":dicomCode / -ContextGroup",
            "skos:exactMatch →", "      SNOMED CT, MeSH"])
em = klass(2, Y0, ":Emission", [":emissionEnergy", ":emissionYield"], [])
rt = klass(2, em[2] + VGAP, ":RadiationType", [], ["skos:notation (ICRP-107 code)"])
row2 = nu[2] + VGAP
db = klass(1, row2, ":DecayBranch", [":branchingRatio", ":decayEnergy"], [])
dm = klass(0, row2, ":DecayMode", [":changeInAtomicNumber", ":changeInMassNumber"], [])

LBL = 338554
mid_gap0 = COL[0] + W + GAP / 2
mid_gap1 = COL[1] + W + GAP / 2
# Element <-> Nuclide
y = Y0 + HDR / 2
arrow(COL[0] + W, y, COL[1], y); label(":hasNuclide", mid_gap0, y - LBL - 20000)
y2 = Y0 + HDR + int(0.42 * IN)
arrow(COL[1], y2, COL[0] + W, y2); label(":isNuclideOf", mid_gap0, y2 + 30000)
# Nuclide -> Emission -> RadiationType
arrow(COL[1] + W, y, COL[2], y); label(":hasEmission", mid_gap1, y - LBL - 20000)
cx = COL[2] + W / 2
arrow(cx, em[2], cx, rt[1]); label(":hasRadiationType", cx - 90000, em[2] + (VGAP - LBL) / 2, "r", int(2.0 * IN))
# Nuclide <-> DecayBranch
cx1 = COL[1] + W / 2 - int(0.42 * IN)
cx2 = COL[1] + W / 2 + int(0.42 * IN)
ly = nu[2] + (VGAP - LBL) / 2
arrow(cx1, nu[2], cx1, row2); label(":hasDecayBranch", cx1 - 90000, ly, "r", int(2.2 * IN))
arrow(cx2, row2, cx2, nu[2]); label(":hasDecayProduct", cx2 + 90000, ly, "l", int(2.0 * IN))
# DecayBranch -> DecayMode
y3 = row2 + HDR / 2
arrow(COL[1], y3, COL[0] + W, y3); label(":hasDecayMode", mid_gap0, y3 - LBL - 20000)

prs.save(OUT)
print("wrote", OUT)
